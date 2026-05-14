from openpyxl import load_workbook
import pandas as pd


def try_repair_excel(file_path):
    try:
        # 尝试用openpyxl打开
        wb = load_workbook(file_path, data_only=True)
        print("✅ openpyxl成功打开文件")

        # 转换为pandas DataFrame
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]

        data = []
        for row in ws.iter_rows(values_only=True):
            data.append(row)

        df = pd.DataFrame(data)
        return df
    except Exception as e:
        print(f"openpyxl打开失败: {e}")
        return None


# 使用
df = try_repair_excel("/home/fxy/project/DAMMFND/weibo21_enhance/train_2_domain.xlsx")
if df is not None:
    print(f"恢复数据成功，形状: {df.shape}")
    print(df.head())